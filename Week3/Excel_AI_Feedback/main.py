import logging
logging.getLogger().setLevel(logging.ERROR)
from google import genai
from google.genai import types
from dotenv import load_dotenv
import os
import openpyxl
import time
load_dotenv()
api_key= os.getenv("GEMINI_API_KEY")
client = genai.Client(api_key=api_key)
SOURCE_FILE = "sample_with_feedback.xlsx" if os.path.exists("sample_with_feedback.xlsx") else "sample.xlsx"
workbook = openpyxl.load_workbook(SOURCE_FILE)
sheet = workbook.active
sheet.cell(row=1, column=5, value= "AI Feedback")
MISSING_MARKS = "Missing marks, feedback not generated."
API_FAILURE= "Something went wrong with feedback generation."
api_calls = 0
for rowNumber, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start = 2):
    name, maths, science, english = row[:4]
    if name is None:
        continue
    if None in (maths, science, english):
        print(f"Missing marks, feedback not generated for {name}")
        sheet.cell(row=rowNumber, column=5, value=MISSING_MARKS)
        workbook.save("sample_with_feedback.xlsx")
        continue
    if (sheet.cell(row=rowNumber, column=5).value is not None) and (sheet.cell(row=rowNumber, column=5).value != MISSING_MARKS) and (sheet.cell(row=rowNumber, column=5).value != API_FAILURE):
        continue
    student_data = f"{name}, Maths {maths}, Science {science}, English {english}"
    response = None

    for attempt in range(3):
        try:
            api_calls +=1
            if api_calls == 14:
                print("Wait for 60 sec.")
                time.sleep(60)
                
            response = client.models.generate_content(
                model="gemini-3.5-flash-lite",
                contents=student_data,
                config=types.GenerateContentConfig(
                    system_instruction="You are a student feedback generator. You will receive a student's name and marks in Mathematics, Science, and English. Generate exactly one inspiring, constructive sentence that acknowledges strong performance where appropriate and identifies subjects that need improvement. Keep the tone encouraging and respectful, and base all feedback strictly on the provided marks. All obtained marks are out of 100. Marks below 50 are considered as fail. Do not appreciate the failed subjects and explicitly consider it as weak subjects.Marks from 50 to 64 are good but they need improvement. If the lowest marks in any subject of a student is not below 65, then do not consider it a weak subject and do not say that they need improvement. Do not make assumptions about the student's intelligence, effort, personality, or background, and do not invent any information. do not quote threshold numbers in the feedback"
                )
            )
            if response is not None and response.text != "" and response.text != None:
                break
        except Exception as e:
            print(f"Attempt: {attempt+1} for {name}. Error: {e}")
            if attempt <2:
                time.sleep((attempt*2) + 2)
    if response and response.text:
        print(student_data+" Feedback: "+response.text)
        sheet.cell(row= rowNumber, column=5, value= response.text)
        workbook.save("sample_with_feedback.xlsx")
    else:
        sheet.cell(row=rowNumber, column= 5, value=API_FAILURE)
        workbook.save("sample_with_feedback.xlsx")
print(f"Total API calls made: {api_calls}")